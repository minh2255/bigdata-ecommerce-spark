"""Build the TV3 report from reviewed Markdown using python-docx.

Run with a Python environment containing python-docx. Black text/borders,
Times New Roman 12 pt throughout, including headings, links and tables.
"""
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT


ROOT = Path(__file__).resolve().parent
FILES = ["ch02_2_spark.md", "ch05_trien_khai.md", "phu_luc_chay_windows.md",
         "ch08_bai_hoc_ai.md", "nguon_chinh_thuc.md"]
OUTPUT = ROOT / "TV3_BaoCao.docx"
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = section.bottom_margin = Cm(1.8)
section.left_margin = section.right_margin = Cm(2.1)
section.footer_distance = Cm(0.8)
for style in doc.styles:
    if style.type in (1, 2):
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().find(qn("w:rFonts"))
        if fonts is not None:
            for key in ["asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"]:
                fonts.attrib.pop(qn("w:" + key), None)
            for key in ["ascii", "hAnsi", "eastAsia", "cs"]:
                fonts.set(qn("w:" + key), "Times New Roman")
    if style.type == 1:
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.08
        style.paragraph_format.widow_control = True
for name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
    doc.styles[name].font.bold = True
    doc.styles[name].paragraph_format.keep_with_next = True
    doc.styles[name].paragraph_format.space_before = Pt(7)
    doc.styles[name].paragraph_format.space_after = Pt(4)
for border in doc.styles.element.xpath("//w:pBdr"):
    border.getparent().remove(border)


def black_run(run, bold=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold = bold
    return run


def inline(paragraph, text):
    for token in re.split(r"(\[[^\]]+\]\([^)]+\)|`[^`]+`|\*\*[^*]+\*\*)", text):
        match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
        if match:
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), paragraph.part.relate_to(match[2], RT.HYPERLINK, is_external=True))
            run = OxmlElement("w:r")
            props = OxmlElement("w:rPr")
            for tag, value in [("color", "000000"), ("sz", "24"), ("szCs", "24")]:
                element = OxmlElement("w:" + tag)
                element.set(qn("w:val"), value)
                props.append(element)
            run.append(props)
            label = OxmlElement("w:t")
            label.text = match[1]
            run.append(label)
            link.append(run)
            paragraph._p.append(link)
        else:
            bold = token.startswith("**") and token.endswith("**")
            token = token[2:-2] if bold else token.strip("`")
            black_run(paragraph.add_run(token), bold=bold or paragraph.style.name.startswith("Heading"))


def table(lines):
    records = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    records = [row for row in records if not all(re.fullmatch(r":?-+:?", cell) for cell in row)]
    grid = doc.add_table(rows=1, cols=len(records[0]))
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid.autofit = False
    available = int(section.page_width - section.left_margin - section.right_margin)
    for cell in grid.rows[0].cells:
        cell.width = available // len(records[0])
    for i, record in enumerate(records):
        row = grid.rows[0] if i == 0 else grid.add_row()
        trpr = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        trpr.append(no_split)
        if i == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for cell, value in zip(row.cells, record):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ["top", "left", "bottom", "right"]:
                element = OxmlElement("w:" + edge)
                for key, val in [("val", "single"), ("sz", "4"), ("color", "000000")]:
                    element.set(qn("w:" + key), val)
                borders.append(element)
            props.append(borders)
            margin = OxmlElement("w:tcMar")
            for edge in ["top", "left", "bottom", "right"]:
                element = OxmlElement("w:" + edge)
                element.set(qn("w:w"), "70")
                element.set(qn("w:type"), "dxa")
                margin.append(element)
            props.append(margin)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            inline(paragraph, value)
            if i == 0:
                for run in paragraph.runs:
                    run.bold = True
    doc.add_paragraph()


doc.add_paragraph("Báo cáo phần việc TV3", "Title")
doc.add_paragraph("Làm sạch dữ liệu thương mại điện tử bằng Apache Spark")
doc.add_paragraph("Ngày 18 tháng 10 năm 2026. Kết quả hiện có là synthetic; dữ liệu Kaggle và xác nhận tích hợp còn chờ.")
for file_index, name in enumerate(FILES):
    if file_index:
        doc.add_page_break()
    lines = (ROOT / name).read_text(encoding="utf-8").splitlines()
    i = 0
    code = False
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            code = not code
            i += 1
            continue
        if line.strip() == "<!-- page -->":
            doc.add_page_break()
        elif code:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(0)
            black_run(paragraph.add_run(line))
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i])
                i += 1
            table(rows)
            continue
        elif line.startswith("#"):
            level = min(len(line) - len(line.lstrip("#")), 3)
            paragraph = doc.add_paragraph(style="Heading " + str(level))
            inline(paragraph, line.lstrip("# "))
        elif line.strip():
            paragraph = doc.add_paragraph()
            inline(paragraph, line)
        i += 1
footer = section.footer.paragraphs[0]
footer.alignment = 2
run = black_run(footer.add_run())
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
run._r.addnext(field)
# Enforce explicit font/color for every run, including content in tables.
for run in doc.element.xpath("//w:r"):
    props = run.find(qn("w:rPr"))
    if props is None:
        props = OxmlElement("w:rPr")
        run.insert(0, props)
    for tag, value in [("color", "000000"), ("sz", "24"), ("szCs", "24")]:
        element = props.find(qn("w:" + tag))
        if element is None:
            element = OxmlElement("w:" + tag)
            props.append(element)
        element.set(qn("w:val"), value)
        element.attrib.pop(qn("w:themeColor"), None)
doc.core_properties.title = "Báo cáo phần việc TV3"
doc.core_properties.subject = "Cleaning và pipeline Spark; báo cáo cá nhân có minh chứng synthetic"
doc.core_properties.author = "TV3"
doc.save(OUTPUT)
print(OUTPUT)
